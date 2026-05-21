"""
===============================================================================
EXCEL UTILS — DATA DRIVEN TESTING (DDTF)
===============================================================================

This utility handles:
    1. Reading test data from Excel (login_test_data.xlsx)
    2. Writing test results back to Excel
    3. Updating date, time, and result status
    4. Applying color formatting (PASS / FAIL)

Used in Data Driven Testing Framework (DDTF)

===============================================================================
"""

import os
from datetime import datetime

from openpyxl import load_workbook
from utils.logger import get_logger


# =============================================================================
# LOGGER SETUP
# =============================================================================

log = get_logger(__name__)


# =============================================================================
# COLUMN MAPPING (EXCEL STRUCTURE)
# =============================================================================

# These numbers represent column positions in Excel sheet (1-based index)
COL_TEST_ID  = 1   # Column A → Test ID
COL_USERNAME = 2   # Column B → Username
COL_PASSWORD = 3   # Column C → Password
COL_DATE     = 4   # Column D → Date
COL_TIME     = 5   # Column E → Time
COL_TESTER   = 6   # Column F → Tester Name
COL_RESULT   = 7   # Column G → Test Result


# =============================================================================
# RESULT CONSTANTS
# =============================================================================

RESULT_PASS = "Test Passed"
RESULT_FAIL = "Test Failed"


# =============================================================================
# EXCEL UTILITY CLASS
# =============================================================================

class ExcelUtils:

    # -------------------------------------------------------------------------
    # DEFAULT EXCEL FILE PATH
    # -------------------------------------------------------------------------

    # Points to: project_root/data/login_test_data.xlsx
    DEFAULT_PATH = os.path.join(
        os.path.dirname(__file__),
        "..",
        "data",
        "login_test_data.xlsx"
    )

    SHEET_NAME = "Login Test Data"

    # =========================================================================
    # METHOD : READ TEST DATA
    # =========================================================================

    # Reads all rows from Excel and converts them into list of dictionaries
    @classmethod
    def read_test_data(cls, filepath: str = None) -> list[dict]:

        # Use default file if no path is provided
        path = filepath or cls.DEFAULT_PATH

        log.info(f"Reading Excel file: {path}")

        # Open Excel workbook
        workbook = load_workbook(path)

        # Select required sheet
        sheet = workbook[cls.SHEET_NAME]

        data = []

        # Start from row 2 (row 1 is header)
        for row in sheet.iter_rows(min_row=2, values_only=False):

            # Skip empty rows
            if all(cell.value is None for cell in row):
                continue

            # Read values from each column
            test_id  = row[COL_TEST_ID - 1].value or ""
            username = row[COL_USERNAME - 1].value or ""
            password = row[COL_PASSWORD - 1].value or ""
            tester   = row[COL_TESTER - 1].value or ""

            # Save actual Excel row number (important for writing back)
            row_num = row[0].row

            # Store row data as dictionary
            data.append({
                "test_id": str(test_id),
                "username": str(username),
                "password": str(password),
                "tester": str(tester),
                "row": row_num
            })

            log.debug(f"Loaded row {row_num} → {test_id}")

        # Close workbook after reading
        workbook.close()

        log.info(f"Total rows loaded: {len(data)}")

        return data

    # =========================================================================
    # METHOD : WRITE TEST RESULT BACK TO EXCEL
    # =========================================================================

    @classmethod
    def write_result(
        cls,
        result: str,
        row_number: int,
        filepath: str = None
    ) -> None:
        """
        This method -
            * Updates:
                - Date
                - Time
                - Result (PASS / FAIL)

            * Also applies color formatting:
                - Green → PASS
                - Red   → FAIL
        """

        # Import styling tools (used only in this method)
        from openpyxl.styles import PatternFill, Font, Alignment

        # Use default file if path not provided
        path = filepath or cls.DEFAULT_PATH

        log.info(f"Writing result '{result}' to row {row_number}")

        # Open workbook
        workbook = load_workbook(path)
        sheet = workbook[cls.SHEET_NAME]

        # Get current date and time
        now = datetime.now()

        # ---------------------------------------------------------------------
        # WRITE DATE AND TIME
        # ---------------------------------------------------------------------

        sheet.cell(row=row_number, column=COL_DATE).value = (
            now.strftime("%Y-%m-%d")
        )

        sheet.cell(row=row_number, column=COL_TIME).value = (
            now.strftime("%H:%M:%S")
        )

        # ---------------------------------------------------------------------
        # WRITE RESULT
        # ---------------------------------------------------------------------

        result_cell = sheet.cell(
            row=row_number,
            column=COL_RESULT
        )

        result_cell.value = result

        # Format text (bold + centered)
        result_cell.font = Font(
            name="Arial",
            size=10,
            bold=True
        )

        result_cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        # ---------------------------------------------------------------------
        # COLOR CODING
        # ---------------------------------------------------------------------

        if result == RESULT_PASS:

            # Green background for PASS
            result_cell.fill = PatternFill(
                "solid",
                start_color="C6EFCE"
            )

            log.info(f"Row {row_number} → PASS")

        else:

            # Red background for FAIL
            result_cell.fill = PatternFill(
                "solid",
                start_color="FFC7CE"
            )

            log.info(f"Row {row_number} → FAIL")

        # Save Excel file
        workbook.save(path)

        # Close workbook
        workbook.close()

        log.info("Excel updated successfully")